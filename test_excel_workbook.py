from __future__ import annotations

import ast
import io
import re
from datetime import date, datetime, timezone
from pathlib import Path
from typing import Any, Optional

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

import sec_nongaap as ng


FUNCTION_NAMES = {"clean_text", "format_date", "excel_ready", "style_data_sheet", "build_bridge_export_frame", "build_excel_export"}


def _load_excel_export_function() -> Any:
    source_path = Path("app.py")
    module = ast.parse(source_path.read_text())
    functions = [node for node in module.body if isinstance(node, ast.FunctionDef) and node.name in FUNCTION_NAMES]
    namespace: dict[str, Any] = {
        "Any": Any,
        "Optional": Optional,
        "date": date,
        "datetime": datetime,
        "timezone": timezone,
        "io": io,
        "re": re,
        "pd": pd,
        "ng": ng,
        "Alignment": Alignment,
        "Border": Border,
        "Font": Font,
        "PatternFill": PatternFill,
        "Side": Side,
        "get_column_letter": get_column_letter,
        "APP_VERSION": ng.APP_VERSION,
    }
    exec(compile(ast.Module(body=functions, type_ignores=[]), str(source_path), "exec"), namespace)
    return namespace["build_excel_export"]


def test_excel_export_contains_primary_sheets_and_clickable_sec_source() -> None:
    build_excel_export = _load_excel_export_function()
    source_url = "https://www.sec.gov/Archives/edgar/data/1/1/earnings.htm"
    reconciliations = pd.DataFrame(
        [
            {
                "pair_id": "pair-1",
                "period": "FY2026 Q2",
                "fiscal_year": 2026,
                "fiscal_quarter": "Q2",
                "period_end": "2026-06-30",
                "metric": "Adjusted EBITDA",
                "gaap_label": "Net income",
                "gaap_display": "$10.0",
                "adjustment_display": "$4.0",
                "non_gaap_label": "Adjusted EBITDA",
                "non_gaap_display": "$14.0",
                "non_gaap_value": 14.0,
                "unit": "usd",
                "scale": "millions",
                "confidence": "High",
                "source_role": "Press release",
                "source_page": 1,
                "source_url": source_url,
            }
        ]
    )
    analysis = {
        "reconciliations": reconciliations,
        "adjustments": pd.DataFrame(),
        "adjustment_history": pd.DataFrame(),
        "adjustment_tieouts": pd.DataFrame(
            [{"pair_id": "pair-1", "tie_out_status": "Ties within rounding", "tie_out_note": "Difference: $0.0"}]
        ),
        "mentions": pd.DataFrame(),
        "definitions": pd.DataFrame(),
        "kpis": pd.DataFrame(),
        "coverage": pd.DataFrame(),
        "sources": pd.DataFrame(),
        "evidence": pd.DataFrame(),
        "request_events": pd.DataFrame(),
        "warnings": pd.DataFrame(),
    }

    workbook_bytes = build_excel_export(
        {"name": "Example Issuer", "ticker": "EXMP", "cik": 1, "sic": "7372", "sic_description": "Software", "fiscal_year_end": "1231"},
        [2026],
        analysis,
        ng.make_metric_matrix(reconciliations),
        reconciliations,
    )

    workbook = load_workbook(io.BytesIO(workbook_bytes))
    assert workbook.sheetnames[:3] == ["Summary", "Extracted metrics", "Reconciliation bridges"]
    metrics_sheet = workbook["Extracted metrics"]
    headers = [cell.value for cell in metrics_sheet[1]]
    assert {"Non-GAAP metric", "Comparable GAAP value", "Reported non-GAAP value", "SEC source"}.issubset(headers)
    source_column = headers.index("SEC source") + 1
    assert metrics_sheet.cell(2, source_column).hyperlink.target == source_url

    bridge_sheet = workbook["Reconciliation bridges"]
    bridge_headers = [cell.value for cell in bridge_sheet[1]]
    assert {"Line item", "Reported value", "SEC source"}.issubset(bridge_headers)
